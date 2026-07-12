"""
Jarvis OS — File content extraction (Phase 6, Part 2)

Turns a file on disk into plain text for the semantic file index. Three
families, one entry point (extract_text):

  .txt/.md/…   plain text — read with the ReadFileTool guards (size cap +
               binary-null detection), decoded utf-8 with errors="replace"
  .pdf         pypdf page text (best-effort — scanned/image PDFs yield little)
  .docx        python-docx paragraph + table text

Everything is BEST-EFFORT and never raises: an unreadable/corrupt/encrypted
file returns None (the indexer skips it), never a crash. Output is capped at
MAX_EXTRACT_CHARS so one huge file cannot blow up memory or the embedder.
Extraction runs off the event loop (extract_text is async → thread pool).
"""
import asyncio
from pathlib import Path
from typing import Optional

from loguru import logger

# ------------------------------------------------------------------ limits
MAX_EXTRACT_CHARS = 200_000   # per file — plenty for search, bounds memory
READ_MAX_BYTES = 5_000_000    # bytes read for a plain-text file
PDF_MAX_PAGES = 200
_BINARY_SNIFF = 8192          # bytes checked for a NUL (binary marker)

# Extension → family. Only these are indexed; everything else is skipped.
TEXT_EXTS = {".txt", ".md", ".markdown", ".text", ".log", ".csv", ".tsv", ".rst"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
INDEXABLE_EXTS = TEXT_EXTS | PDF_EXTS | DOCX_EXTS


def is_indexable(path: Path) -> bool:
    """True when the file's extension is one this module can extract text from."""
    return path.suffix.lower() in INDEXABLE_EXTS


def _cap(text: str) -> str:
    if len(text) <= MAX_EXTRACT_CHARS:
        return text
    return text[:MAX_EXTRACT_CHARS]


def _read_text(path: Path) -> Optional[str]:
    """Plain-text read with the ReadFileTool safety model: size-capped, binary
    files (a NUL in the first block) refused, utf-8 with replacement."""
    try:
        with open(path, "rb") as fh:
            raw = fh.read(READ_MAX_BYTES)
    except OSError as e:
        logger.debug(f"file_extract: cannot read text '{path}': {e}")
        return None
    if b"\x00" in raw[:_BINARY_SNIFF]:
        return None  # looks binary despite a text-ish extension
    text = raw.decode("utf-8", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")  # normalize newlines


def _read_pdf(path: Path) -> Optional[str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            # An empty-password decrypt often works; if not, give up cleanly.
            try:
                reader.decrypt("")
            except Exception:
                return None
        parts: list[str] = []
        total = 0
        for page in reader.pages[:PDF_MAX_PAGES]:
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text:
                parts.append(text)
                total += len(text)
                if total >= MAX_EXTRACT_CHARS:
                    break
        return "\n".join(parts)
    except Exception as e:
        logger.debug(f"file_extract: cannot read pdf '{path}': {e}")
        return None


def _read_docx(path: Path) -> Optional[str]:
    try:
        import docx

        document = docx.Document(str(path))
        parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        logger.debug(f"file_extract: cannot read docx '{path}': {e}")
        return None


def _extract_sync(path: Path) -> Optional[str]:
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        text = _read_text(path)
    elif ext in PDF_EXTS:
        text = _read_pdf(path)
    elif ext in DOCX_EXTS:
        text = _read_docx(path)
    else:
        return None
    if not text or not text.strip():
        return None
    return _cap(text)


async def extract_text(path: Path) -> Optional[str]:
    """Extract plain text from a supported file, or None. Never raises —
    off the event loop, best-effort per the module contract."""
    return await asyncio.to_thread(_extract_sync, path)
